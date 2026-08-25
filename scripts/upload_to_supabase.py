"""
upload_to_supabase.py
Reads the suspect-report Excel and uploads new records to Supabase.

Usage:
  python scripts/upload_to_supabase.py --report <path-to-report.xlsx>
"""
import sys
import subprocess

for pkg in ["openpyxl", "requests"]:
    try:
        __import__(pkg)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", pkg, "--quiet"])

import argparse
import json
import openpyxl
import requests
import uuid
import os
from datetime import datetime, timezone

_NEIGHBORHOODS_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "jerusalem_neighborhoods.json")
with open(_NEIGHBORHOODS_PATH, encoding="utf-8") as _f:
    _NEIGHBORHOOD_REGISTRY = json.load(_f)
_CANONICAL_NEIGHBORHOODS = set(_NEIGHBORHOOD_REGISTRY["neighborhoods"])
_NEIGHBORHOOD_ALIASES = _NEIGHBORHOOD_REGISTRY["aliases"]
_NEIGHBORHOOD_CLEAR_VALUES = set(_NEIGHBORHOOD_REGISTRY["clearValues"])

def normalize_neighborhood(raw):
    v = (raw or "").strip()
    if not v or v in _NEIGHBORHOOD_CLEAR_VALUES:
        return None
    if v in _CANONICAL_NEIGHBORHOODS:
        return v
    return _NEIGHBORHOOD_ALIASES.get(v)

parser = argparse.ArgumentParser()
parser.add_argument("--report", required=True)
args = parser.parse_args()

SUPABASE_URL = "https://mcsygsqfyuaexxxwsgem.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6Im1jc3lnc3FmeXVhZXh4eHdzZ2VtIiwicm9sZSI6ImFub24iLCJpYXQiOjE3Nzg1Njk1MTIsImV4cCI6MjA5NDE0NTUxMn0.Hu6t2PLjE_D113NMQEGvEv8QGhqN6udKNO9McqK3ST8"

HEADERS = {
    "apikey": SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Content-Type": "application/json",
    "Prefer": "return=representation",
}

def supabase_get(table, params=""):
    r = requests.get(f"{SUPABASE_URL}/rest/v1/{table}?{params}", headers=HEADERS)
    r.raise_for_status()
    return r.json()

def supabase_post(table, payload):
    r = requests.post(f"{SUPABASE_URL}/rest/v1/{table}", json=payload, headers=HEADERS)
    if not r.ok:
        raise RuntimeError(f"Supabase error {r.status_code}: {r.text}")
    return r.json()

def map_status(val):
    v = (val or "").strip()
    if v in ("גבוה", "בינוני"):
        return "suspicious"
    if v == "לא חשוד":
        return "ok"
    return "unknown"

def normalize_col(s):
    return " ".join(str(s).split()) if s else ""

COL_ALIASES = {
    "כתובת תואמת":     ["כתובת תואמת (מהעירייה)", "כתובת תואמת"],
    "שמות בעלי נכסים": ["שמות בעלי נכסים באותה כתובת", "שמות בעלי נכסים"],
    "מס' דירות":       ["מס' דירות בכתובת", "מספר דירות בכתובת", "מס' דירות"],
    "מקור/URL":        ["מקור המידע (URL)", "מקור/URL", "מקור המידע"],
}

def resolve_col(col_map, canonical_name):
    if canonical_name in col_map:
        return col_map[canonical_name]
    for alias in COL_ALIASES.get(canonical_name, []):
        if normalize_col(alias) in col_map:
            return col_map[normalize_col(alias)]
    return None

wb = openpyxl.load_workbook(args.report, data_only=True)
ws = wb.active

header_row_idx = None
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    normalized = [normalize_col(v) for v in row]
    if "שם העסק" in normalized:
        header_row_idx = i
        col_map = {normalize_col(v): j for j, v in enumerate(row) if v}
        break

if header_row_idx is None:
    raise ValueError("Header row not found in report Excel")

def col(row_vals, name):
    idx = resolve_col(col_map, name)
    if idx is None:
        return None
    v = row_vals[idx] if idx < len(row_vals) else None
    return str(v).strip() if v is not None else None

records = []
for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
    name = col(row, "שם העסק")
    if not name:
        continue
    records.append({
        "name":                name,
        "type":                col(row, "סוג העסק"),
        "address":             col(row, "כתובת"),
        "neighborhood":        normalize_neighborhood(col(row, "שכונה")),
        "suspicion_rating":    col(row, "דירוג אינדיקציה"),
        "matched_address":     col(row, "כתובת תואמת"),
        "property_owners":     col(row, "שמות בעלי נכסים"),
        "unit_count":          col(row, "מס' דירות"),
        "suspicion_detail":    col(row, "פירוט האינדיקציה"),
        "no_suspicion_reason": col(row, "סיבת אי-אינדיקציה"),
        "link":                col(row, "מקור/URL"),
        "arnona_status":       map_status(col(row, "דירוג אינדיקציה")),
    })

existing_raw = supabase_get("businesses", "select=name,address")
existing_keys = {f"{r['name']}|{r['address']}" for r in existing_raw}

new_records = [r for r in records if f"{r['name']}|{r['address']}" not in existing_keys]
skipped = len(records) - len(new_records)

print(f"Total: {len(records)} | New: {len(new_records)} | Skipped (dup): {skipped}")

if not new_records:
    print("Nothing to upload — all records already exist.")
    sys.exit(0)

file_name = os.path.basename(args.report)
session_id = str(uuid.uuid4())
now_iso = datetime.now(timezone.utc).isoformat()

status_counts = {"suspicious": 0, "ok": 0, "unknown": 0}
for r in new_records:
    status_counts[r["arnona_status"]] = status_counts.get(r["arnona_status"], 0) + 1

supabase_post("upload_sessions", {
    "id":               session_id,
    "file_name":        file_name,
    "upload_date":      now_iso,
    "total_count":      len(new_records),
    "suspicious_count": status_counts["suspicious"],
    "ok_count":         status_counts["ok"],
    "unknown_count":    status_counts["unknown"],
    "skipped_count":    skipped,
    "business_ids":     [],
})

inserted_ids = []
BATCH = 50
for i in range(0, len(new_records), BATCH):
    batch = new_records[i:i + BATCH]
    payload = []
    for r in batch:
        row_id = str(uuid.uuid4())
        inserted_ids.append(row_id)
        payload.append({
            "id":                  row_id,
            "name":                r["name"],
            "type":                r["type"],
            "address":             r["address"],
            "neighborhood":        r["neighborhood"],
            "matched_address":     r["matched_address"],
            "property_owners":     r["property_owners"],
            "unit_count":          r["unit_count"],
            "suspicion_rating":    r["suspicion_rating"],
            "suspicion_detail":    r["suspicion_detail"],
            "no_suspicion_reason": r["no_suspicion_reason"],
            "link":                r["link"],
            "arnona_status":       r["arnona_status"],
            "upload_date":         now_iso,
            "upload_session_id":   session_id,
        })
    supabase_post("businesses", payload)

patch_headers = {**HEADERS, "Prefer": "return=minimal"}
requests.patch(
    f"{SUPABASE_URL}/rest/v1/upload_sessions?id=eq.{session_id}",
    json={"business_ids": inserted_ids},
    headers=patch_headers,
).raise_for_status()

print(f"Uploaded: {len(inserted_ids)} | session_id: {session_id}")
print(f"suspicious: {status_counts['suspicious']} | ok: {status_counts['ok']} | unknown: {status_counts['unknown']}")
print("DONE")
