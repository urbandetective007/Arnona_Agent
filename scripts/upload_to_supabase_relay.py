"""
upload_to_supabase_relay.py
Reads the suspect-report Excel and uploads via Vercel relay function to Supabase.

Usage:
  python scripts/upload_to_supabase_relay.py --report <path-to-report.xlsx> --relay-url <vercel-function-url>
"""
import sys
import subprocess

for pkg in ["openpyxl", "requests"]:
    try:
        __import__(pkg)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", pkg, "--quiet"])

import argparse
import openpyxl
import requests
import uuid
import os
from datetime import datetime, timezone

parser = argparse.ArgumentParser()
parser.add_argument("--report", required=True)
parser.add_argument("--relay-url", required=True)
args = parser.parse_args()

RELAY_URL = args.relay_url.rstrip("/")

def map_status(val):
    v = (val or "").strip()
    if v in ("גבוה", "בינוני"):
        return "suspicious"
    if v == "לא חשוד":
        return "ok"
    return "unknown"

def normalize_col(s):
    return " ".join(str(s).split()) if s else ""

COL_ALIASES = {
    "כתובת תואמת":     ["כתובת תואמת (מהעירייה)", "כתובת תואמת"],
    "שמות בעלי נכסים": ["שמות בעלי נכסים באותה כתובת", "שמות בעלי נכסים"],
    "מס' דירות":       ["מס' דירות בכתובת", "מספר דירות בכתובת", "מס' דירות"],
    "מקור/URL":        ["מקור המידע (URL)", "מקור/URL", "מקור המידע"],
}

def resolve_col(col_map, canonical_name):
    if canonical_name in col_map:
        return col_map[canonical_name]
    for alias in COL_ALIASES.get(canonical_name, []):
        if normalize_col(alias) in col_map:
            return col_map[normalize_col(alias)]
    return None

wb = openpyxl.load_workbook(args.report, data_only=True)
ws = wb.active

header_row_idx = None
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    normalized = [normalize_col(v) for v in row]
    if "שם העסק" in normalized:
        header_row_idx = i
        col_map = {normalize_col(v): j for j, v in enumerate(row) if v}
        break

if header_row_idx is None:
    raise ValueError("Header row not found in report Excel")

def col(row_vals, name):
    idx = resolve_col(col_map, name)
    if idx is None:
        return None
    v = row_vals[idx] if idx < len(row_vals) else None
    return str(v).strip() if v is not None else None

records = []
for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
    name = col(row, "שם העסק")
    if not name:
        continue
    records.append({
        "name":                name,
        "type":                col(row, "סוג העסק"),
        "address":             col(row, "כתובת"),
        "suspicion_rating":    col(row, "דירוג אינדיקציה"),
        "matched_address":     col(row, "כתובת תואמת"),
        "property_owners":     col(row, "שמות בעלי נכסים"),
        "unit_count":          col(row, "מס' דירות"),
        "suspicion_detail":    col(row, "פירוט האינדיקציה"),
        "no_suspicion_reason": col(row, "סיבת אי-אינדיקציה"),
        "link":                col(row, "מקור/URL"),
        "arnona_status":       map_status(col(row, "דירוג אינדיקציה")),
    })

print(f"Total records in report: {len(records)}")

if not records:
    print("No records to upload.")
    sys.exit(0)

# Call relay function to upload
file_name = os.path.basename(args.report)
session_id = str(uuid.uuid4())
now_iso = datetime.now(timezone.utc).isoformat()

status_counts = {"suspicious": 0, "ok": 0, "unknown": 0}
for r in records:
    status_counts[r["arnona_status"]] = status_counts.get(r["arnona_status"], 0) + 1

# Prepare payload with all record details
payload_records = []
for r in records:
    row_id = str(uuid.uuid4())
    payload_records.append({
        "id":                  row_id,
        "name":                r["name"],
        "type":                r["type"],
        "address":             r["address"],
        "matched_address":     r["matched_address"],
        "property_owners":     r["property_owners"],
        "unit_count":          r["unit_count"],
        "suspicion_rating":    r["suspicion_rating"],
        "suspicion_detail":    r["suspicion_detail"],
        "no_suspicion_reason": r["no_suspicion_reason"],
        "link":                r["link"],
        "arnona_status":       r["arnona_status"],
        "upload_date":         now_iso,
    })

# Send to relay function
print(f"Sending {len(payload_records)} records to relay function...")
try:
    relay_response = requests.post(
        f"{RELAY_URL}/api/upload",
        json={
            "session_id": session_id,
            "file_name": file_name,
            "businesses": payload_records,
            "upload_date": now_iso,
            "status_counts": status_counts,
        },
        timeout=120,
    )
    relay_response.raise_for_status()
    result = relay_response.json()

    print(f"Upload via relay successful!")
    print(f"Session ID: {session_id}")
    print(f"Uploaded: {result.get('uploaded', len(payload_records))}")
    print(f"suspicious: {status_counts['suspicious']} | ok: {status_counts['ok']} | unknown: {status_counts['unknown']}")
    print("DONE")

except requests.exceptions.RequestException as e:
    print(f"Error calling relay function: {e}")
    if hasattr(e.response, 'text'):
        print(f"Response: {e.response.text}")
    sys.exit(1)
